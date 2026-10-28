"""Look of the .pptx export (16:9, python-pptx only, fonts Georgia and Calibri).

export.py loads the job's files and calls build_deck(); this file only draws. No model call, no network.
Content rules live elsewhere and are kept here: every bullet is the statement text unchanged, sources sit under each bullet
and in the speaker notes, the Sources slide lists the web sources of ALL statements, the example is shown once.

Fitting: python-pptx does not shrink text, so the code measures it. Text width is estimated with one average character
width per font (conservative, the same on every machine), then lines x line height decide whether the bullets fit the area.
The font size comes from the number of words (20/18/16/14 pt). If the bullets still do not fit, the LAST bullets move to the
speaker notes (with their sources) and a warning is added to `warnings`. Nothing is allowed to leave the slide.
"""
import re
from urllib.parse import urlparse

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from schemas import Section, Slide, Statement
from slides import source_text

NAVY, TEAL, AMBER = "0F2A4A", "0E9F8E", "F59E0B"
BG, INK, MUTED = "F7F8FA", "1F2933", "64748B"
CHIP_FILL, BOX_FILL, BOX_LINE = "E3F4F1", "FFFFFF", "D5DCE5"
ON_NAVY, ON_NAVY_MUTED = "FFFFFF", "AEBBCB"
SERIF, SANS = "Georgia", "Calibri"

SLIDE_W, SLIDE_H = 13.333, 7.5
LEFT, RIGHT = 0.6, 12.733
BODY_TOP, BODY_BOTTOM = 1.7, 6.85
CHAR_W = {SERIF: 0.60, SANS: 0.45}  # average character width in em; PowerPoint renders Calibri text at about 0.42, Georgia bold at about 0.58
LINE = 1.2  # line height in em
MIN_SIZE = 14
GAP_JOIN = 2.0  # lecture ranges that are closer than this many seconds are shown as one range (display only)
HYPHEN = "\u2011"  # a non-breaking hyphen keeps "01:01-01:10" on one line
SOURCE_SIZE = 12
TITLE_CHARS = 60  # a web title longer than this is shortened on the slide (the link keeps the exact url)


EXAMPLE_LEAD = re.compile(r"^\s*for\s+(example|instance)\s*[,:]\s*", re.IGNORECASE)


def example_display(text: str) -> str:
    """For display only: the label is "Example:", so a leading "For example," / "For instance," is removed. notes.json is not changed."""
    text = " ".join(text.split())
    stripped = EXAMPLE_LEAD.sub("", text, count=1)
    return stripped[:1].upper() + stripped[1:] if stripped != text and stripped else text


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


# --------------------------------------------------------------------------------------------- measuring
def lines_needed(text: str, size_pt: float, width_in: float, font: str = SANS, bold: bool = False) -> int:
    """Greedy word wrap with one average character width. Words longer than a line count as several lines."""
    char_in = CHAR_W[font] * size_pt * (1.07 if bold else 1.0) / 72
    per_line = max(1, int(width_in / char_in))
    lines, used = 1, 0
    for word in text.split():
        n = len(word)
        if used and used + 1 + n > per_line:
            lines, used = lines + 1, 0
        if n > per_line:
            lines += (n - 1) // per_line
            used = n % per_line or per_line
        else:
            used = used + (1 if used else 0) + n
    return lines


def height_in(lines: int, size_pt: float) -> float:
    return lines * size_pt * LINE / 72


def body_size(words: int) -> int:
    return 20 if words <= 80 else 18 if words <= 110 else 16 if words <= 140 else MIN_SIZE


def clock(seconds: float) -> str:
    total = int(seconds)
    if total >= 3600:
        return f"{total // 3600}:{total % 3600 // 60:02d}:{total % 60:02d}"
    return f"{total // 60:02d}:{total % 60:02d}"


def short(text: str, limit: int) -> str:
    text = " ".join(text.split())
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def domain(url: str) -> str:
    host = urlparse(url).netloc.lower()
    return host[4:] if host.startswith("www.") else host


def lecture_ranges(statement: Statement) -> list[tuple[float, float]]:
    """The lecture ranges of a statement as they are SHOWN: sorted, touching or nearly touching ranges joined."""
    return merge_ranges([(s.start, s.end) for s in statement.sources if s.type == "lecture"])


def fmt_range(a: float, b: float) -> str:
    return f"{clock(a)}{HYPHEN}{clock(b)}"


def web_sources(statement: Statement) -> list[tuple[str, str]]:
    """[(label, exact url)]: "Title (domain)"; the label is only what is displayed."""
    out = []
    for s in statement.sources:
        if s.type == "web":
            label = short(s.title, TITLE_CHARS) if s.title else ""
            out.append((f"{label} ({domain(s.url)})" if label else domain(s.url) or s.url, s.url))
    return out


def source_plain(statement: Statement, tag: str | None) -> str:
    """The source line as plain text, used to measure it."""
    parts = []
    ranges = lecture_ranges(statement)
    if ranges:
        parts.append("Lecture " + ", ".join(fmt_range(a, b) for a, b in ranges))
    parts += [label for label, _ in web_sources(statement)]
    line = " | ".join(parts)
    return f"{line}   {tag}" if tag else line


# ---------------------------------------------------------------------------------------------- drawing
def _background(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _run(paragraph, text, size, color, font=SANS, bold=False, italic=False, link=None):
    r = paragraph.add_run()
    r.text = text  # plain text only: python-pptx escapes it, nothing from the data is interpreted
    r.font.size = Pt(size)
    r.font.name = font
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    if link:
        r.hyperlink.address = link
    return r


def _frame(shape, margin=0.05, anchor=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    return tf


def _textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return box, _frame(box, anchor=anchor)


def _bullet_format(paragraph, color: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.3)))
    pPr.set("indent", str(-Inches(0.3)))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = pPr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": color})
    clr.append(srgb)
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def _indent_plain(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.3)))
    pPr.set("indent", "0")
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def _slide_number(slide, n: int, color: str) -> None:
    box, tf = _textbox(slide, 11.9, 7.0, 0.85, 0.3)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _run(p, str(n), 11, color)


def _title(slide, text: str, x, y, w, h, color, max_lines, sizes, anchor=MSO_ANCHOR.TOP):
    """Use the layout's real title placeholder (outline view, screen readers) and style it."""
    size = next((s for s in sizes if lines_needed(text, s, w - 0.2, SERIF, True) <= max_lines), sizes[-1])
    t = slide.shapes.title
    lines = lines_needed(text, size, w - 0.2, SERIF, True)
    t.left, t.top, t.width, t.height = Inches(x), Inches(y), Inches(w), Inches(min(h, 0.8) if lines == 1 and max_lines <= 2 else h)
    tf = _frame(t, margin=0.05, anchor=anchor)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    _run(p, text, size, color, SERIF, bold=True)
    return size, lines


# ------------------------------------------------------------------------------------------- slide kinds
def _title_slide(prs, title, stats_line, number):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    _background(s, NAVY)
    _rect(s, 0, 0, 0.18, SLIDE_H, TEAL)
    _title(s, title, 0.9, 1.4, 11.5, 2.0, ON_NAVY, 3, (44, 38, 32, 28), anchor=MSO_ANCHOR.BOTTOM)
    _rect(s, 0.95, 3.65, 1.4, 0.06, TEAL)
    box, tf = _textbox(s, 0.9, 3.9, 11.5, 0.6)
    _run(tf.paragraphs[0], stats_line, 20, ON_NAVY)
    box, tf = _textbox(s, 0.9, 6.55, 10.8, 0.5)
    _run(tf.paragraphs[0], "Built with Nemotron on Nebius Token Factory. Every statement cites a lecture time or a web link.", 13, ON_NAVY_MUTED)
    _slide_number(s, number, ON_NAVY_MUTED)
    return s


def _tag_color(tag: str) -> str:
    return TEAL if tag in ("verified", "rewritten") else MUTED


def _fit(bullets: list[Statement], example: str | None, tags: dict, area_h: float, size: int, width: float) -> int:
    """How many of the bullets (from the first) fit in area_h inches at this size."""
    used, count = 0.1, 0  # text frame margins
    for i, b in enumerate(bullets):
        h = height_in(lines_needed(b.text, size, width - 0.3), size)
        h += height_in(lines_needed(source_plain(b, tags.get(b.id)), SOURCE_SIZE, width - 0.3), SOURCE_SIZE) + 3 / 72
        h += 0 if i == 0 else 10 / 72
        if used + h > area_h:
            break
        used, count = used + h, count + 1
    return count


def _content_slide(prs, slide: Slide, time_range: str | None, tags: dict, number: int, warnings: list[str]):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    _background(s, BG)
    _rect(s, 0, 0, 0.18, SLIDE_H, TEAL)
    _, title_lines = _title(s, slide.title, LEFT, 0.4, 9.4, 1.15, NAVY, 2, (32, 28, 24))
    body_top = BODY_TOP if title_lines > 1 else BODY_TOP - 0.4  # a one-line title leaves more room for the bullets
    if time_range:
        chip = _rect(s, 10.23, 0.52, RIGHT - 10.23, 0.42, CHIP_FILL, MSO_SHAPE.ROUNDED_RECTANGLE)
        chip.adjustments[0] = 0.5
        tf = _frame(chip, margin=0.05, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, f"Lecture {time_range}", 13, TEAL, bold=True)

    width = RIGHT - LEFT
    bullets = list(slide.bullets)
    words = sum(len(b.text.split()) for b in bullets) + len((slide.example or "").split())
    size = body_size(words)
    notes_extra: list[str] = []

    # the example box is placed first (bottom of the body); the bullets get what is left above it
    example = slide.example
    box_h = 0.0
    if example:
        text = example_display(example)
        ex_size = max(MIN_SIZE, size - 2)
        max_h = 1.9
        while True:
            box_h = 0.2 + 0.28 + height_in(lines_needed(text, ex_size, width - 0.6), ex_size) + 0.12
            if box_h <= max_h or ex_size == MIN_SIZE:
                break
            ex_size -= 1
        if box_h > max_h:  # last resort: shorten what is DISPLAYED; the full text goes to the speaker notes
            keep = len(text)
            while box_h > max_h and keep > 20:
                keep = int(keep * 0.85)
                shown = short(text, keep)
                box_h = 0.2 + 0.28 + height_in(lines_needed(shown, ex_size, width - 0.6), ex_size) + 0.12
            notes_extra.append(f"Example (full text): {text}")
            warnings.append(f'slide "{slide.title}": the example was shortened on the slide; the full text is in the speaker notes')
            text = shown
        area_bottom = BODY_BOTTOM - box_h - 0.2
    else:
        area_bottom = BODY_BOTTOM
    area_h = area_bottom - body_top

    count = _fit(bullets, example, tags, area_h, size, width)
    if count == 0 and bullets:  # one statement alone is too long even for an empty slide: show its beginning
        first = bullets[0]
        keep = len(first.text)
        while count == 0 and keep > 20:
            keep = int(keep * 0.85)
            probe = first.model_copy(update={"text": short(first.text, keep)})
            count = _fit([probe], example, tags, area_h, size, width)
        bullets[0] = probe
        notes_extra.append(f"{first.text} [Source: {source_text(first)}]")
        warnings.append(f'slide "{slide.title}": statement {first.id} was shortened on the slide; the full text is in the speaker notes')
        count = 1
    moved = bullets[count:]
    bullets = bullets[:count]
    if moved:
        notes_extra = [f"{m.text} [Source: {source_text(m)}]" for m in moved] + notes_extra
        warnings.append(f'slide "{slide.title}": {len(moved)} bullet(s) did not fit and were moved to the speaker notes: '
                        + ", ".join(m.id for m in moved))

    box, tf = _textbox(s, LEFT, body_top, width, area_h)
    first_par = True
    for b in bullets:
        p = tf.paragraphs[0] if first_par else tf.add_paragraph()
        if not first_par:
            p.space_before = Pt(10)
        first_par = False
        _bullet_format(p, TEAL)
        _run(p, " ".join(b.text.split()), size, INK)
        q = tf.add_paragraph()
        q.space_before = Pt(3)
        _indent_plain(q)
        parts_written = False
        ranges = lecture_ranges(b)
        if ranges:
            _run(q, "Lecture " + ", ".join(fmt_range(a, z) for a, z in ranges), SOURCE_SIZE, MUTED)
            parts_written = True
        for label, url in web_sources(b):
            if parts_written:
                _run(q, " | ", SOURCE_SIZE, MUTED)
            _run(q, label, SOURCE_SIZE, TEAL, link=url)
            parts_written = True
        tag = tags.get(b.id)
        if tag:
            _run(q, ("   " if parts_written else "") + tag, SOURCE_SIZE, _tag_color(tag), bold=True)

    if example:
        y = BODY_BOTTOM - box_h
        card = _rect(s, LEFT, y, width, box_h, BOX_FILL, MSO_SHAPE.ROUNDED_RECTANGLE)
        card.adjustments[0] = 0.12
        card.line.fill.solid()
        card.line.fill.fore_color.rgb = rgb(BOX_LINE)
        card.line.width = Pt(1)
        tf = _frame(card, margin=0.2, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        _run(p, "EXAMPLE", 12, AMBER, bold=True)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.space_before = Pt(2)
        _run(p2, text, ex_size, INK)

    notes = slide.speaker_notes
    if notes_extra:
        notes = "\n".join(notes_extra + ([notes] if notes else []))
    s.notes_slide.notes_text_frame.text = notes
    _slide_number(s, number, MUTED)
    return s


def merge_ranges(ranges: list[tuple[float, float]], gap: float = GAP_JOIN) -> list[tuple[float, float]]:
    out: list[list[float]] = []
    for a, b in sorted(ranges):
        if out and a <= out[-1][1] + gap:
            out[-1][1] = max(out[-1][1], b)
        else:
            out.append([a, b])
    return [(a, b) for a, b in out]


def _paginate(items: list[float], budget: float) -> list[list[int]]:
    """Group item indexes into pages whose heights (inches) fit the budget."""
    pages, current, used = [], [], 0.0
    for i, h in enumerate(items):
        if current and used + h > budget:
            pages.append(current)
            current, used = [], 0.0
        current.append(i)
        used += h
    return pages + [current] if current else [[]]


def _sources_slides(prs, slides: list[Slide], statements: list[Statement], number: int) -> int:
    col_w = (RIGHT - LEFT - 0.5) / 2
    text_w = col_w - 0.1
    size = 13
    by_section: dict[str, list[tuple[float, float]]] = {}
    web: dict[str, str] = {}
    for st in statements:
        by_section.setdefault(st.section_id, []).extend(lecture_ranges(st))
        for s in st.sources:
            if s.type == "web":
                web.setdefault(s.url, s.title or s.url)

    left = []  # (section title, ranges text)
    for sl in slides:
        merged = merge_ranges(by_section.get(sl.section_id, []))
        if merged:
            left.append((sl.title, ", ".join(fmt_range(a, b) for a, b in merged)))
    right = [(f"{short(title, 80)} ({domain(url)})", url) for url, title in web.items()]

    left_h = [height_in(lines_needed(t, size, text_w, SANS, True) + lines_needed(r, size, text_w), size) + 0.14 for t, r in left]
    right_h = [height_in(lines_needed(label, size, text_w), size) + 0.14 for label, _ in right]
    budget = BODY_BOTTOM - BODY_TOP - 0.7
    left_pages, right_pages = _paginate(left_h, budget), _paginate(right_h, budget)
    total = max(len(left_pages), len(right_pages))
    for page in range(total):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        _background(s, BG)
        _rect(s, 0, 0, 0.18, SLIDE_H, TEAL)
        _title(s, "Sources" + (f" ({page + 1}/{total})" if total > 1 else ""), LEFT, 0.4, 9.4, 1.15, NAVY, 2, (32, 28, 24))
        for col, heading in enumerate(("In the lecture", "On the web")):
            x = LEFT + col * (col_w + 0.5)
            box, tf = _textbox(s, x, BODY_TOP, col_w, BODY_BOTTOM - BODY_TOP)
            p = tf.paragraphs[0]
            _run(p, heading, 20, TEAL, SERIF, bold=True)
            if col == 0:
                for idx in (left_pages[page] if page < len(left_pages) else []):
                    title, ranges = left[idx]
                    h = tf.add_paragraph()
                    h.space_before = Pt(8)
                    _run(h, title, size, INK, bold=True)
                    r = tf.add_paragraph()
                    _run(r, ranges, size, MUTED)
            else:
                shown = right_pages[page] if page < len(right_pages) else []
                if not right and page == 0:
                    q = tf.add_paragraph()
                    q.space_before = Pt(8)
                    _run(q, "No web sources were used. Every point comes from the lecture (time ranges are under each bullet).", size, MUTED)
                for idx in shown:
                    label, url = right[idx]
                    q = tf.add_paragraph()
                    q.space_before = Pt(8)
                    _run(q, label, size, TEAL, link=url)
        _slide_number(s, number + page, MUTED)
    return number + total


# ------------------------------------------------------------------------------------------------ public
def _theme_link_colors(prs) -> None:
    """Hyperlinks take their colour from the theme (a run colour is ignored): set it to the teal of the palette."""
    import re as _re

    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    theme = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")
    xml = _re.sub(r'(<a:hlink>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', r"\g<1>" + TEAL, xml)
    xml = _re.sub(r'(<a:folHlink>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', r"\g<1>" + "0B7A6E", xml)
    theme._blob = xml.encode("utf-8")


def build_deck(
    slides: list[Slide], *, title: str = "Lecture notes", statements: list[Statement] | None = None,
    sections: list[Section] | None = None, status: dict[str, str] | None = None, judged: bool | None = None,
    warnings: list[str] | None = None,
) -> bytes:
    """status: {statement_id: "verified" | "rewritten" | "not judged"} from the judge files, or None when there is no data.
    judged False = the job has no notes_final.json: the title slide says "not judged" and every statement is tagged so."""
    import io

    warnings = warnings if warnings is not None else []
    statements = statements if statements is not None else [b for sl in slides for b in sl.bullets]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    _theme_link_colors(prs)
    prs.core_properties.title = title
    if judged is False:
        status = {st.id: "not judged" for st in statements}
    status = status or {}

    section_by_id = {sec.id: sec for sec in sections or []}
    n_sections = len(sections) if sections else len(slides)
    stats = [f"{n_sections} section{'s' if n_sections != 1 else ''}", f"{len(statements)} statement{'s' if len(statements) != 1 else ''}"]
    if judged is False:
        stats.append("not judged")
    elif status:
        stats.append(f"{sum(status.get(st.id) in ('verified', 'rewritten') for st in statements)} verified by the judge")
    _title_slide(prs, title, " - ".join(stats), 1)

    number = 2
    for sl in slides:
        sec = section_by_id.get(sl.section_id)
        if sec is not None:
            time_range = f"{clock(sec.start)}-{clock(sec.end)}"
        else:  # no sections given: the span of the lecture sources of this section's statements
            spans = [r for st in statements if st.section_id == sl.section_id for r in lecture_ranges(st)] or \
                    [r for b in sl.bullets for r in lecture_ranges(b)]
            time_range = f"{clock(min(a for a, _ in spans))}-{clock(max(b for _, b in spans))}" if spans else None
        _content_slide(prs, sl, time_range, status, number, warnings)
        number += 1
    _sources_slides(prs, slides, statements, number)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
