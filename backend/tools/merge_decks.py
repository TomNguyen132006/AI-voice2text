"""Merge several finished jobs (lectures) into ONE .pptx and ONE notes.md. Local only: no model call, no network.

Usage (from backend/):
    python -m tools.merge_decks --manifest lectures.json --out OUTDIR [--title "Lecture notes"]
lectures.json (kept by the user, NOT in git): [{"title": "Introduction", "job": "job_..."}, ...] in the order they should appear.

The drawing is done by deck.py (unchanged): the slides of all lectures are given to deck.build_deck at once, then a divider slide is
inserted before each lecture and the slide numbers are written again. Every lecture has a MODE, taken from its files:
    full:  the job went through the judge (notes_final.json exists)        -> "full: source-checked"
    light: PIPELINE_MODE=light, or no notes_final.json                      -> "light: not judged"
A statement of a light lecture is never tagged "source-checked", in the deck or in the notes. Ids are prefixed ("L3_st_0001") because
every job numbers its statements and sections from 1. Timestamps and web links are kept as they are.
"""
import argparse
import io
import json
import sys
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

import deck
import export
from schemas import Section, Slide, Statement
from storage import get_storage, job_key, load_json
from topics import load_sections

MODE_LABELS = {"full": "full: source-checked", "light": "light: not judged"}


@dataclass
class Lecture:
    n: int
    title: str
    job_id: str
    mode: str  # "full" | "light"
    sections: list[Section]
    statements: list[Statement]
    slides: list[Slide]
    status: dict[str, str]  # statement id -> "source-checked" | "rewritten" | "not judged" (ids already prefixed)

    @property
    def tag(self) -> str:
        return f"L{self.n}_"


def _optional(job_id, name, storage):
    try:
        return load_json(job_key(job_id, name), storage)
    except FileNotFoundError:
        return None


def detect_mode(job_id: str, storage=None) -> tuple[str, bool]:
    """(mode, judged). light when the pipeline report says so, or when the judge never produced notes_final.json."""
    storage = storage or get_storage()
    judged = _optional(job_id, "notes_final.json", storage) is not None
    report = _optional(job_id, "pipeline_report.json", storage) or {}
    light = report.get("mode") == "light" or not judged
    return ("light" if light else "full"), judged and not light


def load_lecture(n: int, title: str, job_id: str, storage=None) -> Lecture:
    storage = storage or get_storage()
    if _optional(job_id, "slides.json", storage) is None:
        raise FileNotFoundError(f"{job_id}: slides.json is missing, so this job has no slides to merge")
    mode, judged = detect_mode(job_id, storage)
    raw = _optional(job_id, "notes_final.json" if judged else "notes.json", storage)
    if raw is None:
        raise FileNotFoundError(f"{job_id}: neither notes_final.json nor notes.json exists")
    statements = [Statement(**s) for s in raw]
    slides = [Slide(**s) for s in _optional(job_id, "slides.json", storage)]
    sections = load_sections(job_id, storage)
    status = {st.id: "not judged" for st in statements}
    if mode == "full":
        found = export._judge_status(job_id, statements, storage)  # None when judge.json is missing: then nothing is claimed
        if found:
            status = found
    tag = f"L{n}_"
    return Lecture(
        n=n, title=title, job_id=job_id, mode=mode,
        sections=[s.model_copy(update={"id": tag + s.id}) for s in sections],
        statements=[s.model_copy(update={"id": tag + s.id, "section_id": tag + s.section_id}) for s in statements],
        slides=[
            sl.model_copy(update={"section_id": tag + sl.section_id, "bullets": [b.model_copy(update={"id": tag + b.id, "section_id": tag + b.section_id}) for b in sl.bullets]})
            for sl in slides
        ],
        status={tag + k: v for k, v in status.items()},
    )


def load_lectures(manifest: list[dict], storage=None) -> list[Lecture]:
    return [load_lecture(i, item["title"], item["job"], storage) for i, item in enumerate(manifest, start=1)]


# ----------------------------------------------------------------------------------------------- the deck
def _divider(prs, lecture: Lecture, total: int):
    """Navy slide before each lecture: its number, title, mode, and what is in it. Built with deck.py's own helpers."""
    s = prs.slides.add_slide(prs.slide_layouts[5])
    deck._background(s, deck.NAVY)
    deck._rect(s, 0, 0, 0.18, deck.SLIDE_H, deck.TEAL)
    box, tf = deck._textbox(s, 0.9, 1.7, 11.5, 0.5)
    deck._run(tf.paragraphs[0], f"LECTURE {lecture.n} OF {total}", 14, deck.ON_NAVY_MUTED, bold=True)
    deck._title(s, lecture.title, 0.9, 2.2, 11.5, 1.8, deck.ON_NAVY, 3, (44, 38, 32, 28), anchor=deck.MSO_ANCHOR.TOP)
    deck._rect(s, 0.95, 4.2, 1.4, 0.06, deck.TEAL)
    box, tf = deck._textbox(s, 0.9, 4.45, 11.5, 0.6)
    deck._run(tf.paragraphs[0], MODE_LABELS[lecture.mode], 22, deck.ON_NAVY, bold=True)
    box, tf = deck._textbox(s, 0.9, 5.1, 11.5, 0.6)
    counts = f"{len(lecture.sections)} section{'s' if len(lecture.sections) != 1 else ''} - {len(lecture.statements)} statements"
    note = "Every statement was checked against its cited source by the judge." if lecture.mode == "full" else "These notes did not go through the judge."
    deck._run(tf.paragraphs[0], f"{counts}. {note}", 16, deck.ON_NAVY_MUTED)
    deck._slide_number(s, 0, deck.ON_NAVY_MUTED)  # the real number is written by _renumber
    return s


def _renumber(prs) -> None:
    """deck.py writes the slide number as plain text; after dividers are inserted the numbers must be written again."""
    for number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.left == Inches(11.9) and shape.top == Inches(7.0):
                runs = shape.text_frame.paragraphs[0].runs
                if runs:
                    runs[0].text = str(number)


def merge_deck(lectures: list[Lecture], title: str = "Lecture notes", warnings: list[str] | None = None) -> bytes:
    warnings = warnings if warnings is not None else []
    slides = [sl for lec in lectures for sl in lec.slides]
    statements = [st for lec in lectures for st in lec.statements]
    sections = [sec for lec in lectures for sec in lec.sections]
    status = {k: v for lec in lectures for k, v in lec.status.items()}
    data = deck.build_deck(slides, title=title, statements=statements, sections=sections, status=status, judged=True, warnings=warnings)
    prs = Presentation(io.BytesIO(data))

    ids = list(prs.slides._sldIdLst)  # order: title, all content slides, Sources pages
    content = ids[1 : 1 + len(slides)]
    tail = ids[1 + len(slides) :]
    dividers = [_divider(prs, lec, len(lectures)) for lec in lectures]
    divider_ids = list(prs.slides._sldIdLst)[len(ids) :]

    order, at = [ids[0]], 0
    for lec, divider in zip(lectures, divider_ids):
        order.append(divider)
        order += content[at : at + len(lec.slides)]
        at += len(lec.slides)
    order += tail
    lst = prs.slides._sldIdLst
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    _renumber(prs)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------------------------- the notes
def merge_notes(lectures: list[Lecture], title: str = "Lecture notes") -> str:
    lines = [f"# {title}", "", "Contents:", ""]
    lines += [f"{lec.n}. {lec.title} ({MODE_LABELS[lec.mode]})" for lec in lectures]
    lines += ["", "Lectures marked \"light: not judged\" did not go through the judge: a point may not be supported by its source.", ""]
    for lec in lectures:
        body = export.notes_md(lec.sections, lec.statements, title=f"Lecture {lec.n}: {lec.title}", judged=lec.mode == "full")
        lines += ["---", "", f"_Mode: {MODE_LABELS[lec.mode]}_", "", body.strip(), ""]
    return "\n".join(lines).rstrip() + "\n"


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Merge finished jobs into one deck and one notes.md (local, no model call).")
    parser.add_argument("--manifest", required=True, help='JSON list [{"title": ..., "job": ...}] in lecture order')
    parser.add_argument("--out", required=True, help="folder for merged.pptx and merged_notes.md")
    parser.add_argument("--title", default="Lecture notes")
    args = parser.parse_args(argv)
    manifest = json.loads(Path(args.manifest).read_text(encoding="utf-8"))
    lectures = load_lectures(manifest)
    warnings: list[str] = []
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    (out / "merged.pptx").write_bytes(merge_deck(lectures, args.title, warnings))
    (out / "merged_notes.md").write_text(merge_notes(lectures, args.title), encoding="utf-8")
    for lec in lectures:
        print(f"{lec.n:2d}. {lec.title} | {MODE_LABELS[lec.mode]} | {len(lec.sections)} sections, {len(lec.statements)} statements | {lec.job_id}")
    for w in warnings:
        print("WARNING:", w, file=sys.stderr)
    print(f"written to {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
