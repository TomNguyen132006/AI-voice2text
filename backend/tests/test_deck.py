"""Tests for deck.py (the look of the .pptx). No network, no keys, no model."""
import io
import json

import pytest
from pptx import Presentation
from pptx.util import Inches

import deck
import export
from schemas import Section, Slide, Statement
from storage import LocalStorage, job_key, save_json

HYPHEN = deck.HYPHEN


def lec(start, end, seg="seg_0001"):
    return {"type": "lecture", "segment_id": seg, "start": start, "end": end}


def web(n, title=None):
    return {"type": "web", "url": f"https://www.site{n}.example/path?q={n}&x=1", "title": title or f"Page {n}"}


def st(n, text=None, sources=None, section="sec_01"):
    return Statement(id=f"st_{n:04d}", section_id=section, text=text or f"Statement number {n} is short.",
                     sources=sources if sources is not None else [lec(n * 10, n * 10 + 5)])


def make(bullets, example=None, notes="Lecture time of this section: 00:00 - 01:00 (to listen again).", title="Topic", section="sec_01"):
    return Slide(section_id=section, title=title, bullets=bullets, example=example, speaker_notes=notes)


def build(slides, **kw):
    warnings = kw.pop("warnings", [])
    data = deck.build_deck(slides, warnings=warnings, **kw)
    return Presentation(io.BytesIO(data)), warnings


def slide_text(slide):
    return "\n".join(p.text for sh in slide.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs)


def bullet_texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                pPr = p._p.pPr
                if pPr is not None and pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None:
                    out.append(p.text)
    return out


def links(slide):
    return [r.hyperlink.address for sh in slide.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs
            for r in p.runs if r.hyperlink.address]


def notes(slide):
    return slide.notes_slide.notes_text_frame.text


SECTIONS = [Section(id="sec_01", title="Topic", start=12.0, end=95.0, segment_ids=["seg_0001"])]


# ------------------------------------------------------------------------------------------ structure
def test_slide_is_16_by_9_and_slide_count_is_sections_plus_title_plus_sources():
    slides = [make([st(1)], section="sec_01"), make([st(2, section="sec_02")], title="Second", section="sec_02")]
    prs, _ = build(slides)
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.001
    assert prs.slide_width == Inches(deck.SLIDE_W) and prs.slide_height == Inches(deck.SLIDE_H)
    assert len(prs.slides) == 2 + 2
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["Lecture notes", "Topic", "Second", "Sources"]  # real title placeholders on every slide


def test_every_bullet_is_exactly_its_statement_text():
    bullets = [st(1, "A  database\nis a set of related data."), st(2, "Second one: x < y & z > 1 <b>bold</b>."), st(3)]
    prs, w = build([make(bullets)])
    assert bullet_texts(prs.slides[1]) == [" ".join(b.text.split()) for b in bullets]
    assert w == []
    # user text is plain text, never markup
    assert "<b>bold</b>" in slide_text(prs.slides[1])


def test_each_web_source_has_a_hyperlink_with_the_exact_url_and_title_with_domain_is_shown():
    b1 = st(1, sources=[lec(12, 27), web(1), web(2, "Second page")])
    b2 = st(2, sources=[web(3)])
    prs, _ = build([make([b1, b2])])
    content = prs.slides[1]
    assert links(content) == [web(1)["url"], web(2)["url"], web(3)["url"]]  # exact, query string included
    text = slide_text(content)
    assert "Page 1 (site1.example)" in text and "Second page (site2.example)" in text  # "www." is removed from the displayed domain
    assert f"Lecture 00:12{HYPHEN}00:27 | Page 1 (site1.example) | Second page (site2.example)" in text


def test_example_is_shown_once_with_its_label_and_without_a_doubled_for_example():
    prs, _ = build([make([st(1), st(2)], example="For example, a table has rows.")])
    text = slide_text(prs.slides[1])
    assert text.count("EXAMPLE") == 1 and text.count("A table has rows.") == 1
    assert "For example" not in text
    prs, _ = build([make([st(1)])])
    assert "EXAMPLE" not in slide_text(prs.slides[1])


def test_source_line_has_ranges_joined_when_they_touch_and_a_status_tag():
    b = st(1, sources=[lec(31, 35), lec(35, 39), lec(40, 44), lec(100, 110)])
    prs, _ = build([make([b])], status={"st_0001": "rewritten"}, judged=True)
    text = slide_text(prs.slides[1])
    assert f"Lecture 00:31{HYPHEN}00:44, 01:40{HYPHEN}01:50   rewritten" in text
    assert deck.merge_ranges([(0, 5), (5, 9), (11.5, 12), (30, 31)]) == [(0, 9), (11.5, 12), (30, 31)]


def test_chip_shows_the_section_range_with_or_without_sections():
    b = st(1, sources=[lec(20, 30)]), st(2, sources=[lec(50, 70)])
    prs, _ = build([make(list(b))], sections=SECTIONS)
    assert "Lecture 00:12-01:35" in slide_text(prs.slides[1])
    prs, _ = build([make(list(b))])  # no sections: the span of the lecture sources
    assert "Lecture 00:20-01:10" in slide_text(prs.slides[1])


def test_no_shape_extends_beyond_the_slide_in_any_case():
    long_text = " ".join(f"word{i}" for i in range(220))
    cases = [
        [make([st(1)], example="An example.")],
        [make([st(n) for n in range(1, 8)], example="An example of moderate length that explains the idea a little.")],
        [make([st(n) for n in range(1, 8)])],
        [make([st(1, long_text), st(2, long_text)], example=long_text)],
        [make([st(1)], title="A very long slide title that keeps going and going and going until it needs to be shortened a lot")],
        [make([st(n, sources=[lec(n, n + 2), web(n)]) for n in range(1, 8)], example="Short.")],
    ]
    for slides in cases:
        prs, _ = build(slides, title="A long lecture title " * 6)
        for number, slide in enumerate(prs.slides, start=1):
            for sh in slide.shapes:
                assert sh.left >= 0 and sh.top >= 0, (number, sh.name)
                assert sh.left + sh.width <= prs.slide_width, (number, sh.name)
                assert sh.top + sh.height <= prs.slide_height, (number, sh.name)


# --------------------------------------------------------------------------------------------- fitting
def test_font_size_follows_the_number_of_words():
    assert [deck.body_size(w) for w in (1, 80, 81, 110, 111, 140, 141, 500)] == [20, 20, 18, 18, 16, 16, 14, 14]
    prs, _ = build([make([st(1)])])
    sizes = {r.font.size.pt for sh in prs.slides[1].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs
             if p._p.pPr is not None and p._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None}
    assert sizes == {20}


def test_one_statement_slide_fits_and_shows_it():
    prs, w = build([make([st(1)])])
    assert bullet_texts(prs.slides[1]) == [st(1).text] and w == []


def test_seven_statements_either_fit_or_the_last_ones_go_to_the_notes_with_their_sources_and_a_warning():
    bullets = [st(n, sources=[lec(n * 10, n * 10 + 5), web(n)]) for n in range(1, 8)]
    slide = make(bullets, example="Some example text that takes a line or two on the slide.", notes="Lecture time of this section: 00:00 - 01:00 (to listen again).")
    prs, warnings = build([slide])
    shown = bullet_texts(prs.slides[1])
    assert shown == [b.text for b in bullets[: len(shown)]]  # always the FIRST ones, in order
    moved = bullets[len(shown):]
    if moved:
        assert len(warnings) == 1 and ", ".join(m.id for m in moved) in warnings[0]
        for m in moved:
            assert f"{m.text} [Source: " in notes(prs.slides[1]) and m.sources[1].url in notes(prs.slides[1])
    else:
        assert warnings == []
    assert len(shown) + len(moved) == 7
    assert notes(prs.slides[1]).endswith("(to listen again).")  # the original notes are kept


def test_seven_short_statements_without_example_all_fit():
    prs, w = build([make([st(n, f"Short note {n}.") for n in range(1, 8)])])
    assert len(bullet_texts(prs.slides[1])) == 7 and w == []


def test_very_long_statements_never_leave_the_slide_and_nothing_is_lost():
    long_text = " ".join(f"word{i}" for i in range(220))
    slide = make([st(1, long_text), st(2, "Another statement.")], example=long_text)
    prs, warnings = build([slide])
    n = notes(prs.slides[1])
    assert long_text in n  # the full texts are in the speaker notes
    assert warnings and all(w.startswith('slide "Topic"') for w in warnings)
    assert bullet_texts(prs.slides[1])  # the slide still shows something


def test_speaker_notes_keep_the_slide_notes_with_sources_untouched_when_everything_fits():
    text = "Another note. [Source: lecture 00:10-00:15 (seg_0002)]\nLecture time of this section: 00:00 - 01:00 (to listen again)."
    prs, _ = build([make([st(1)], notes=text)])
    assert notes(prs.slides[1]) == text


# -------------------------------------------------------------------------------------------- title slide
def test_title_slide_numbers():
    statements = [st(1), st(2), st(3)]
    status = {"st_0001": "source-checked", "st_0002": "rewritten", "st_0003": "not judged"}
    prs, _ = build([make(statements)], title="Databases 101", statements=statements, sections=SECTIONS, status=status, judged=True)
    text = slide_text(prs.slides[0])
    assert "Databases 101" in text
    assert "1 section - 3 statements - 2 source-checked by the judge" in text
    assert "Built with Nemotron on Nebius Token Factory. Every statement cites a lecture time or a web link." in text
    assert prs.slides[0].shapes.title.text == "Databases 101"
    assert "Lecturesque" in text and prs.core_properties.author == "Lecturesque"


def test_title_slide_says_not_judged_for_a_job_without_judge_and_has_no_made_up_number_without_data():
    statements = [st(1), st(2)]
    prs, _ = build([make(statements)], statements=statements, judged=False)
    assert "1 section - 2 statements - not judged" in slide_text(prs.slides[0])
    assert slide_text(prs.slides[1]).count("not judged") == 2  # every statement is tagged
    prs, _ = build([make(statements)], statements=statements, judged=True, status=None)  # judged, but no per-statement data
    text = slide_text(prs.slides[0])
    assert "1 section - 2 statements" in text and "source-checked" not in text and "not judged" not in text
    assert "source-checked" not in slide_text(prs.slides[1])


# ------------------------------------------------------------------------------------------ sources slide
def test_sources_slide_has_two_columns_and_all_web_sources_deduplicated():
    shown = st(1, sources=[lec(12, 27), web(1)])
    hidden = st(2, sources=[web(2), lec(40, 50)])  # only in the speaker notes
    again = st(3, sources=[web(1)])
    prs, _ = build([make([shown])], statements=[shown, hidden, again])
    last = prs.slides[-1]
    text = slide_text(last)
    assert "In the lecture" in text and "On the web" in text
    assert "Page 1 (site1.example)" in text and "Page 2 (site2.example)" in text
    assert text.count("Page 1") == 1  # deduplicated by url
    assert links(last) == [web(1)["url"], web(2)["url"]]
    assert "No web sources were used" not in text
    assert f"00:12{HYPHEN}00:27, 00:40{HYPHEN}00:50" in text  # time ranges of the section, all statements


def test_no_web_sources_sentence_only_when_the_job_has_none():
    prs, _ = build([make([st(1)])], statements=[st(1)])
    assert "No web sources were used" in slide_text(prs.slides[-1])
    prs, _ = build([make([st(1)])], statements=[st(1), st(2, sources=[web(9)])])
    assert "No web sources were used" not in slide_text(prs.slides[-1])


def test_sources_slide_is_split_in_pages_when_long():
    many = [st(n, sources=[web(n, f"A fairly long page title number {n} about databases and their models")]) for n in range(1, 40)]
    prs, _ = build([make(many[:2])], statements=many)
    titles = [s.shapes.title.text for s in prs.slides]
    pages = [t for t in titles if t.startswith("Sources")]
    assert len(pages) > 1 and pages == [f"Sources ({i}/{len(pages)})" for i in range(1, len(pages) + 1)]
    all_links = [u for s in list(prs.slides)[-len(pages):] for u in links(s)]
    assert all_links == [web(n)["url"] for n in range(1, 40)]


def test_slide_numbers_are_on_every_slide():
    prs, _ = build([make([st(1)]), make([st(2)], title="Two")], statements=[st(1), st(2)])
    for n, slide in enumerate(prs.slides, start=1):
        assert slide_text(slide).strip().splitlines()[-1] == str(n)


# --------------------------------------------------------------------------------------------- export wiring
def test_export_job_status_tags_come_from_the_judge_files(tmp_path):
    storage = LocalStorage(tmp_path)
    statements = [st(1), st(2), st(3), st(4)]
    save_json(job_key("job_d1", "sections.json"), [s.model_dump() for s in SECTIONS], storage)
    save_json(job_key("job_d1", "notes_final.json"), [s.model_dump() for s in statements], storage)
    save_json(job_key("job_d1", "slides.json"), [make(statements).model_dump()], storage)
    save_json(job_key("job_d1", "transcript_clean.json"), [{"id": "seg_0001", "start": 0, "end": 5, "text": "x"}], storage)
    save_json(job_key("job_d1", "judge.json"), [{"statement_id": "st_0001", "verdict": "pass", "reason": "", "by": "ultra"},
                                               {"statement_id": "st_0002", "verdict": "unsupported", "reason": "", "by": "ultra"},
                                               {"statement_id": "st_0003", "verdict": "unjudged", "reason": "", "by": "none"}], storage)
    save_json(job_key("job_d1", "fix_log.json"), [{"statement_id": "st_0002", "outcome": "fixed"}], storage)
    warnings = []
    export.export_job("job_d1", storage, warnings=warnings)
    prs = Presentation(io.BytesIO(storage.load_bytes(job_key("job_d1", "slides.pptx"))))
    text = slide_text(prs.slides[1])
    assert text.count("source-checked") == 1 and text.count("rewritten") == 1 and text.count("not judged") == 2  # st_0003 and st_0004 (no verdict)
    assert "1 section - 4 statements - 2 source-checked by the judge" in slide_text(prs.slides[0])


def test_export_job_without_judge_files_has_no_tags_and_without_notes_final_says_not_judged(tmp_path):
    storage = LocalStorage(tmp_path)
    statements = [st(1)]
    for name, data in [("sections.json", [s.model_dump() for s in SECTIONS]), ("notes_final.json", [s.model_dump() for s in statements]),
                       ("slides.json", [make(statements).model_dump()]), ("transcript_clean.json", [{"id": "seg_0001", "start": 0, "end": 5, "text": "x"}])]:
        save_json(job_key("job_d2", name), data, storage)
    export.export_job("job_d2", storage)
    prs = Presentation(io.BytesIO(storage.load_bytes(job_key("job_d2", "slides.pptx"))))
    assert "source-checked" not in slide_text(prs.slides[1]) and "not judged" not in slide_text(prs.slides[0])
    assert export._judge_status("job_d2", statements, storage) is None


# ---------------------------------------------------------------------------------- polish: contrast, layout
def _lum(hex_color):
    channels = [int(hex_color[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    channels = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in channels]
    return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]


def contrast(a, b):
    hi, lo = sorted((_lum(a), _lum(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


@pytest.mark.parametrize("name, fg, bg", [
    ("body text", deck.INK, deck.BG), ("body text on white", deck.INK, "FFFFFF"), ("title", deck.NAVY, deck.BG),
    ("source line and slide number", deck.MUTED, deck.BG), ("muted on white", deck.MUTED, "FFFFFF"),
    ("tags, links, headings (dark teal)", deck.TEAL_TEXT, deck.BG), ("chip text", deck.TEAL_TEXT, deck.CHIP_FILL),
    ("EXAMPLE label", deck.AMBER_TEXT, deck.BOX_FILL), ("title slide text", deck.ON_NAVY, deck.NAVY),
    ("title slide footer and number", deck.ON_NAVY_MUTED, deck.NAVY),
])
def test_small_text_has_at_least_4_5_to_1_contrast(name, fg, bg):
    assert contrast(fg, bg) >= 4.5, (name, round(contrast(fg, bg), 2))


def test_the_label_uses_the_dark_amber_and_the_light_amber_is_only_the_border():
    prs, _ = build([make([st(1)], example="An example.")])
    label_runs = [r for sh in prs.slides[1].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.text == "EXAMPLE"]
    assert len(label_runs) == 1 and str(label_runs[0].font.color.rgb) == deck.AMBER_TEXT
    cards = [sh for sh in prs.slides[1].shapes if sh.has_text_frame and "EXAMPLE" in sh.text_frame.text]
    assert str(cards[0].line.color.rgb) == deck.AMBER


def _example_card_and_bullets(slide):
    card = next(sh for sh in slide.shapes if sh.has_text_frame and "EXAMPLE" in sh.text_frame.text)
    box = next(sh for sh in slide.shapes if sh.has_text_frame and sh is not card and sh.text_frame.text.startswith("Statement"))
    return card, box


def test_example_box_sits_right_below_the_bullets_and_short_content_leaves_the_space_at_the_bottom():
    prs, _ = build([make([st(1)], example="An example.")])
    card, box = _example_card_and_bullets(prs.slides[1])
    gap = (card.top - (box.top + box.height)) / 914400
    assert abs(gap - deck.EXAMPLE_GAP) < 0.01  # a fixed gap below the last bullet
    assert card.top < Inches(3.5)  # not pinned to the bottom: the empty band is under the box
    assert card.top + card.height < Inches(5.0)


def test_example_box_never_overlaps_the_bullets_and_stays_inside_the_slide_when_the_slide_is_full():
    bullets = [st(n, f"Statement {n} " + "word " * 12) for n in range(1, 8)]
    prs, _ = build([make(bullets, example="An example with some words in it.")])
    card, box = _example_card_and_bullets(prs.slides[1])
    assert card.top >= box.top + box.height
    assert card.top + card.height <= Inches(deck.BODY_BOTTOM) + 1  # inside the body area
    assert card.top + card.height <= prs.slide_height
