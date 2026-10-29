"""Tests for pipeline.py (S7-7). Fake model client, fake search, no network, no keys."""
import io
import json
import zipfile

import pytest
from pptx import Presentation

import clean
import enrich
import fix
import gaps
import judge
import llm
import notes
import pipeline
import slides
import topics
from storage import LocalStorage, job_key, load_json
from tests.test_clean import FakeClient, payload_of

SRT = """1
00:00:00,000 --> 00:00:10,000
Today we talk about gradient descent.

2
00:00:10,000 --> 00:00:25,000
The gradient tells us which way the loss goes up fastest.

3
00:00:25,000 --> 00:00:40,000
So we take a step the other way, and the loss goes down.

4
00:00:40,000 --> 00:00:55,000
The size of the step is the learning rate.
"""


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    monkeypatch.setattr(llm, "_sleep", lambda s: None)
    monkeypatch.setenv("LLM_LOG_PATH", str(tmp_path / "calls.jsonl"))
    monkeypatch.setenv("SEARCH_CACHE_DIR", str(tmp_path / "search_cache"))


def router(kwargs, n):
    """A fake model that answers every stage in the format that stage expects."""
    system, payload = kwargs["messages"][0]["content"], payload_of(kwargs)
    if system.startswith(clean.SYSTEM_PROMPT[:40]):
        return json.dumps([{"id": s["id"], "text": s["text"]} for s in payload["segments"]])
    if system.startswith(topics.SYSTEM_PROMPT[:40]):
        return json.dumps([{"title": "Gradient descent", "first_segment_id": payload["segments"][0]["id"]}])
    if system.startswith(gaps.SYSTEM_PROMPT[:40]):
        return json.dumps([{"section_id": payload["section_id"], "term": "gradient", "reason": "Needs explaining.", "query": "gradient descent basics"}])
    if system.startswith(notes.SYSTEM_PROMPT[:40]):
        return json.dumps([
            {"text": "The gradient shows the direction in which the loss grows fastest.", "sources": [payload["segments"][1]["id"], "web_1"]},
            {"text": "The learning rate sets the size of each step.", "sources": [payload["segments"][3]["id"]]},
        ])
    if system.startswith(judge.SYSTEM_PROMPT[:40]):
        return json.dumps([{"statement_id": s["id"], "verdict": "pass", "reason": "stated"} for s in payload["statements"]])
    if system.startswith(slides.SYSTEM_PROMPT[:40]):
        return json.dumps({"title": "Gradient and steps", "example": None})
    raise AssertionError(f"unknown stage prompt: {system[:60]}")


def start(tmp_path, name="lecture.srt", data=SRT.encode("utf-8")):
    storage = LocalStorage(tmp_path / "data")
    pipeline.create_job("job_p1", name, data, storage)
    return storage


def test_full_run_on_a_transcript_upload(tmp_path):
    storage = start(tmp_path)
    client = FakeClient(router)
    job = pipeline.run_pipeline("job_p1", "lecture.srt", storage, client=client, provider=enrich.FakeProvider())
    assert job.status == "done" and job.stage == "done" and job.error is None
    assert job.outputs.txt == "/jobs/job_p1/files/transcript.txt"
    for name in ["transcript.json", "transcript_clean.json", "sections.json", "gaps.json", "search.json", "notes.json",
                 "judge.json", "notes_final.json", "slides.json", "transcript.txt", "notes.md", "slides.pptx", "pipeline_report.json", "call_stats.json"]:
        storage.load_bytes(job_key("job_p1", name))  # raises FileNotFoundError if a stage forgot to save
    assert pipeline.load_job("job_p1", storage).status == "done"
    final = load_json(job_key("job_p1", "notes_final.json"), storage)
    assert len(final) == 2 and all(s["sources"] for s in final)
    md = storage.load_bytes(job_key("job_p1", "notes.md")).decode("utf-8")
    assert "📍" in md and "🌐" in md and "learning rate" in md
    titles = [s.shapes.title.text for s in Presentation(io.BytesIO(storage.load_bytes(job_key("job_p1", "slides.pptx")))).slides]
    assert titles[1] == "Gradient and steps"
    report = load_json(job_key("job_p1", "pipeline_report.json"), storage)
    assert set(report["seconds_per_stage"]) == set(pipeline.STAGES) and report["model_calls"]["total_calls"] > 0
    stats_saved = load_json(job_key("job_p1", "call_stats.json"), storage)
    assert stats_saved["by_stage"]["judge"]["calls"] == 1 and stats_saved["total"]["calls"] == report["model_calls"]["total_calls"]
    # only ONE uploaded copy plus derived files: the upload is kept for re-runs
    assert storage.load_bytes(pipeline.upload_key("job_p1", "lecture.srt")) == SRT.encode("utf-8")


def test_stage_order_and_progress(tmp_path, monkeypatch):
    storage = start(tmp_path)
    seen = []
    for name in ["_transcribe_stage", "_clean_stage", "_enrich_stage", "_write_stage", "_judge_stage", "_export_stage"]:
        monkeypatch.setattr(pipeline, name, lambda *a, _n=name, **k: seen.append((_n, pipeline.load_job("job_p1", storage).stage)))
    job = pipeline.run_pipeline("job_p1", "lecture.srt", storage)
    assert [stage for _, stage in seen] == pipeline.STAGES  # job.json shows the stage while it runs
    assert job.status == "done"


def test_a_failing_stage_becomes_an_error_job(tmp_path, monkeypatch):
    storage = start(tmp_path)
    ran = []
    monkeypatch.setattr(pipeline, "_transcribe_stage", lambda *a: ran.append("t"))
    monkeypatch.setattr(pipeline, "_clean_stage", lambda *a: (_ for _ in ()).throw(RuntimeError("model is down")))
    monkeypatch.setattr(pipeline, "_enrich_stage", lambda *a: ran.append("e"))
    job = pipeline.run_pipeline("job_p1", "lecture.srt", storage)
    assert (job.status, job.stage) == ("error", "cleaning") and job.error == "cleaning: model is down"
    assert ran == ["t"]  # later stages did not run
    assert pipeline.load_job("job_p1", storage).error == "cleaning: model is down"


def test_error_message_hides_keys_and_is_short(tmp_path, monkeypatch):
    storage = start(tmp_path)
    monkeypatch.setenv("NEBIUS_API_KEY", "super-secret-key-123456")
    monkeypatch.setattr(pipeline, "_transcribe_stage", lambda *a: (_ for _ in ()).throw(RuntimeError("bad super-secret-key-123456 " + "x" * 1000)))
    job = pipeline.run_pipeline("job_p1", "lecture.srt", storage)
    assert "super-secret" not in job.error and len(job.error) <= pipeline.ERROR_MAX_CHARS


def test_web_search_failure_is_only_a_warning(tmp_path, monkeypatch):
    storage = start(tmp_path)

    def broken_search(*a, **k):
        raise RuntimeError("search is down")

    monkeypatch.setattr(enrich, "search_for_gaps", broken_search)
    job = pipeline.run_pipeline("job_p1", "lecture.srt", storage, client=FakeClient(router))
    assert job.status == "done"
    assert load_json(job_key("job_p1", "search.json"), storage) == []
    assert "web search skipped" in load_json(job_key("job_p1", "pipeline_report.json"), storage)["warnings"][0]
    assert "](http" not in storage.load_bytes(job_key("job_p1", "notes.md")).decode("utf-8")  # no web link


def test_empty_transcript_is_an_error(tmp_path):
    storage = start(tmp_path, "empty.srt", b"\n")
    job = pipeline.run_pipeline("job_p1", "empty.srt", storage, client=FakeClient(router))
    assert job.status == "error" and job.stage == "transcribing"


def test_audio_upload_goes_to_the_transcriber(tmp_path, monkeypatch):
    import transcribe
    from schemas import Segment

    storage = start(tmp_path, "lecture.mp3", b"not really audio")
    got = {}

    def fake_transcribe(path, engine=None):
        got["suffix"], got["bytes"] = path.suffix, path.read_bytes()
        return [Segment(id="seg_0001", start=0, end=5, text="hello")]

    monkeypatch.setattr(transcribe, "transcribe", fake_transcribe)
    monkeypatch.setattr(pipeline, "_clean_stage", lambda *a: (_ for _ in ()).throw(RuntimeError("stop here")))
    job = pipeline.run_pipeline("job_p1", "lecture.mp3", storage)
    assert got == {"suffix": ".mp3", "bytes": b"not really audio"} and job.stage == "cleaning"
    assert load_json(job_key("job_p1", "transcript.json"), storage)[0]["text"] == "hello"


# --- light mode -----------------------------------------------------------------------------------------------

class NoSearch:
    """A search provider that fails the test if light mode ever calls it."""

    def search(self, *args, **kwargs):
        raise AssertionError("light mode must not search the web")


def system_prompts(client):
    return [c["messages"][0]["content"] for c in client.calls]


def run_light(tmp_path, monkeypatch, provider=None):
    monkeypatch.setenv("PIPELINE_MODE", "light")
    storage = start(tmp_path)
    client = FakeClient(router)
    job = pipeline.run_pipeline("job_p1", "lecture.srt", storage, client=client, provider=provider or NoSearch())
    return storage, client, job


def test_light_mode_never_calls_gaps_search_judge_or_fixer(tmp_path, monkeypatch):
    storage, client, job = run_light(tmp_path, monkeypatch)
    assert job.status == "done"
    used = system_prompts(client)
    for forbidden in (gaps.SYSTEM_PROMPT, judge.SYSTEM_PROMPT, fix.SYSTEM_PROMPT):
        assert not any(p.startswith(forbidden[:40]) for p in used)
    for needed in (clean.SYSTEM_PROMPT, topics.SYSTEM_PROMPT, notes.SYSTEM_PROMPT, slides.SYSTEM_PROMPT):
        assert any(p.startswith(needed[:40]) for p in used)
    report = load_json(job_key("job_p1", "pipeline_report.json"), storage)
    assert report["mode"] == "light" and set(report["seconds_per_stage"]) == {"transcribing", "cleaning", "writing", "exporting"}
    by_stage = load_json(job_key("job_p1", "call_stats.json"), storage)["by_stage"]
    assert not {"judge", "fixer", "gap_finder"} & set(by_stage) and {"cleaning", "topic_split", "notes_writer", "slide_writer"} <= set(by_stage)


def test_light_mode_writes_no_judge_files_and_still_makes_all_three_outputs(tmp_path, monkeypatch):
    storage, _, job = run_light(tmp_path, monkeypatch)
    for name in ("notes_final.json", "judge.json", "judge_report.json", "fix_log.json", "gaps.json", "search.json"):
        with pytest.raises(FileNotFoundError):
            storage.load_bytes(job_key("job_p1", name))
    for name in ("notes.json", "slides.json", "notes.md", "slides.pptx", "transcript.txt"):
        storage.load_bytes(job_key("job_p1", name))
    assert job.outputs.md and job.outputs.pptx and job.outputs.txt


def test_light_mode_notes_and_deck_say_not_judged(tmp_path, monkeypatch):
    storage, _, _ = run_light(tmp_path, monkeypatch)
    md = storage.load_bytes(job_key("job_p1", "notes.md")).decode("utf-8")
    assert "NOT JUDGED" in md and "source-checked" not in md.lower()
    prs = Presentation(io.BytesIO(storage.load_bytes(job_key("job_p1", "slides.pptx"))))
    text = lambda slide: "\n".join(p.text for sh in slide.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs)
    assert "not judged" in text(prs.slides[0]) and "source-checked" not in text(prs.slides[0])
    assert "not judged" in text(prs.slides[1]) and "source-checked" not in text(prs.slides[1])


def test_light_mode_slides_use_notes_json(tmp_path, monkeypatch):
    storage, _, _ = run_light(tmp_path, monkeypatch)
    drafted = [s["id"] for s in load_json(job_key("job_p1", "notes.json"), storage)]
    shown = [b["id"] for slide in load_json(job_key("job_p1", "slides.json"), storage) for b in slide["bullets"]]
    assert drafted and set(shown) <= set(drafted)


def test_full_mode_is_the_default_and_unchanged(tmp_path, monkeypatch):
    monkeypatch.delenv("PIPELINE_MODE", raising=False)
    storage = start(tmp_path)
    client = FakeClient(router)
    job = pipeline.run_pipeline("job_p1", "lecture.srt", storage, client=client, provider=enrich.FakeProvider())
    assert job.status == "done"
    used = system_prompts(client)
    for needed in (gaps.SYSTEM_PROMPT, judge.SYSTEM_PROMPT):
        assert any(p.startswith(needed[:40]) for p in used)
    report = load_json(job_key("job_p1", "pipeline_report.json"), storage)
    assert report["mode"] == "full" and set(report["seconds_per_stage"]) == set(pipeline.STAGES)
    storage.load_bytes(job_key("job_p1", "notes_final.json"))


def test_explicit_full_mode_behaves_like_the_default(tmp_path, monkeypatch):
    monkeypatch.setenv("PIPELINE_MODE", " FULL ")
    assert pipeline.pipeline_mode() == "full"
    storage = start(tmp_path)
    job = pipeline.run_pipeline("job_p1", "lecture.srt", storage, client=FakeClient(router), provider=enrich.FakeProvider())
    assert job.status == "done" and set(load_json(job_key("job_p1", "pipeline_report.json"), storage)["seconds_per_stage"]) == set(pipeline.STAGES)


def test_unknown_mode_is_a_clear_error_and_makes_no_model_call(tmp_path, monkeypatch):
    monkeypatch.setenv("PIPELINE_MODE", "fast")
    storage = start(tmp_path)
    client = FakeClient(router)
    job = pipeline.run_pipeline("job_p1", "lecture.srt", storage, client=client)
    assert job.status == "error" and "PIPELINE_MODE must be one of full, light" in job.error and "'fast'" in job.error
    assert client.calls == []
    with pytest.raises(ValueError):
        pipeline.pipeline_mode()
