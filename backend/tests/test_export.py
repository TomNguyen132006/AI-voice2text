"""Tests for export.py (Epic 6). No network, no keys."""
import io
import json
import zipfile
from pathlib import Path

from pptx import Presentation

import export
from schemas import Section, Segment, Slide, Statement
from storage import LocalStorage, job_key, save_json, save_transcript

SAMPLES = Path(export.__file__).parent / "samples"


def sample(name):
    return json.loads((SAMPLES / name).read_text(encoding="utf-8"))


SEGMENTS = [Segment(**s) for s in sample("transcript.json")]
SECTIONS = [Section(**s) for s in sample("sections.json")]
STATEMENTS = [Statement(**s) for s in sample("statements.json")]
SLIDES = [Slide(**s) for s in sample("slides.json")]


def test_clock():
    assert export.clock(75.5) == "01:15" and export.clock(0) == "00:00" and export.clock(3723) == "1:02:03"


def test_transcript_txt_has_one_stamped_line_per_segment():
    text = export.transcript_txt(SEGMENTS)
    lines = text.strip().split("\n")
    assert len(lines) == len(SEGMENTS)
    assert lines[1].startswith("[00:12] The loss function")


def test_notes_md_has_sources_on_every_point():
    md = export.notes_md(SECTIONS, STATEMENTS)
    assert md.startswith("# Lecture notes") and "## Loss and gradient" in md
    bullets = [b for b in md.split("\n- ")[1:]]
    assert len(bullets) == len(STATEMENTS)
    assert all("📍" in b for b in bullets)  # every sample statement has a lecture source
    assert "🌐 [Gradient descent - Wikipedia](https://en.wikipedia.org/wiki/Gradient_descent)" in md
    assert "📍 12:" not in md and "📍 00:12-00:27" in md


def test_notes_md_skips_sections_without_statements_and_handles_none():
    md = export.notes_md(SECTIONS, [s for s in STATEMENTS if s.section_id == "sec_01"])
    assert "Choosing the learning rate" not in md
    assert "No statement passed the checks" in export.notes_md(SECTIONS, [])


def test_markdown_brackets_in_titles_do_not_break_links():
    st = Statement(id="st_1", section_id="sec_01", text="x", sources=[{"type": "web", "url": "https://a.example/p", "title": "A [b] c"}])
    assert "[A (b) c](https://a.example/p)" in export.source_markdown(st)


def test_zip_holds_all_files():
    data = export.build_zip({"a.txt": b"A", "b.md": b"B"})
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert sorted(z.namelist()) == ["a.txt", "b.md"] and z.read("a.txt") == b"A"


def test_export_job_writes_three_files(tmp_path):
    storage = LocalStorage(tmp_path)
    save_transcript("job_x1", SEGMENTS, storage)
    save_json(job_key("job_x1", "sections.json"), [s.model_dump() for s in SECTIONS], storage)
    save_json(job_key("job_x1", "notes.json"), [s.model_dump() for s in STATEMENTS], storage)
    save_json(job_key("job_x1", "slides.json"), [s.model_dump() for s in SLIDES], storage)
    keys = export.export_job("job_x1", storage)
    assert set(keys) == {"txt", "md", "pptx"}
    assert storage.load_bytes(keys["txt"]).decode("utf-8").startswith("[00:00] Today we talk")
    assert b"## Loss and gradient" in storage.load_bytes(keys["md"])
    assert storage.load_bytes(keys["pptx"])[:2] == b"PK"  # a pptx is a zip
    # notes_final.json wins over notes.json
    save_json(job_key("job_x1", "notes_final.json"), [STATEMENTS[0].model_dump()], storage)
    export.export_job("job_x1", storage)
    assert storage.load_bytes(keys["md"]).decode("utf-8").count("\n- ") == 1


def test_notes_md_marks_notes_that_were_not_judged():
    judged = export.notes_md(SECTIONS, STATEMENTS)
    assert "NOT JUDGED" not in judged
    unjudged = export.notes_md(SECTIONS, STATEMENTS, judged=False)
    assert unjudged.splitlines()[2].startswith("**NOT JUDGED:**")  # right under the title, before the first section
    assert unjudged.index("NOT JUDGED") < unjudged.index("## ")
    assert "NOT JUDGED" in export.notes_md(SECTIONS, [], judged=False) and "No statement passed" in export.notes_md(SECTIONS, [], judged=False)


def test_export_job_without_notes_final_warns_and_marks_the_md(tmp_path, capsys):
    storage = LocalStorage(tmp_path)
    save_transcript("job_x2", SEGMENTS, storage)
    save_json(job_key("job_x2", "sections.json"), [s.model_dump() for s in SECTIONS], storage)
    save_json(job_key("job_x2", "notes.json"), [s.model_dump() for s in STATEMENTS], storage)
    save_json(job_key("job_x2", "slides.json"), [s.model_dump() for s in SLIDES], storage)
    keys = export.export_job("job_x2", storage)
    assert "NOT JUDGED" in capsys.readouterr().err
    assert "NOT JUDGED" in storage.load_bytes(keys["md"]).decode("utf-8")
    save_json(job_key("job_x2", "notes_final.json"), [STATEMENTS[0].model_dump()], storage)
    export.export_job("job_x2", storage)
    assert capsys.readouterr().err == ""
    assert "NOT JUDGED" not in storage.load_bytes(keys["md"]).decode("utf-8")


def _web_statement(i, section="sec_01"):
    return Statement(id=f"st_w{i}", section_id=section, text="x",
                     sources=[{"type": "web", "url": f"https://e.example/{i}", "title": f"Page {i}"}])


def _sources_text(data):
    return Presentation(io.BytesIO(data)).slides[-1].placeholders[1].text_frame.text


def test_example_label_is_not_doubled():
    assert export.example_display("For example, a table has rows.") == "A table has rows."
    assert export.example_display("for instance: the 2222 row.") == "The 2222 row."
    assert export.example_display("  For  example,   x is 1. ") == "X is 1."
    assert export.example_display("A table has rows.") == "A table has rows."
    assert export.example_display("For examples we use tables.") == "For examples we use tables."  # only the lead-in phrase
