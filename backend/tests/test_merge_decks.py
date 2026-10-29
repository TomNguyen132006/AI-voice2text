"""Tests for tools/merge_decks.py: one deck and one notes.md from several finished jobs. No network, no model."""
import io

import pytest
from pptx import Presentation

from schemas import Slide, Statement
from storage import LocalStorage, job_key, save_json
from tools import merge_decks as md

LEC = {"type": "lecture", "segment_id": "seg_0001", "start": 12.0, "end": 27.0}
WEB = {"type": "web", "url": "https://www.example.org/page?x=1", "title": "A page"}


def st(n, text, sources, section="sec_01"):
    return Statement(id=f"st_{n:04d}", section_id=section, text=text, sources=sources)


def make_job(storage, job_id, title, statements, mode, judge=None, judge_files=True):
    """Every job numbers its sections and statements from 1, so two jobs always collide on the same ids."""
    sections = [{"id": "sec_01", "title": title, "start": 0.0, "end": 100.0, "segment_ids": ["seg_0001"]}]
    slide = Slide(section_id="sec_01", title=title, bullets=statements, example=None, speaker_notes="n")
    save_json(job_key(job_id, "sections.json"), sections, storage)
    save_json(job_key(job_id, "slides.json"), [slide.model_dump()], storage)
    if mode == "full":
        save_json(job_key(job_id, "notes_final.json"), [s.model_dump() for s in statements], storage)
        if judge_files:
            save_json(job_key(job_id, "judge.json"), [{"statement_id": s.id, "verdict": "pass", "reason": "", "by": "ultra"} for s in statements], storage)
            save_json(job_key(job_id, "fix_log.json"), judge or [], storage)
    else:
        save_json(job_key(job_id, "notes.json"), [s.model_dump() for s in statements], storage)
        save_json(job_key(job_id, "pipeline_report.json"), {"mode": "light"}, storage)


@pytest.fixture
def storage(tmp_path, monkeypatch):
    monkeypatch.setenv("STORAGE_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("STORAGE_BACKEND", "local")
    s = LocalStorage(tmp_path / "data")
    make_job(s, "job_full", "Full lecture", [st(1, "Full statement one.", [LEC]), st(2, "Full web statement.", [WEB])], "full",
             judge=[{"statement_id": "st_0002", "outcome": "fixed"}])
    make_job(s, "job_light", "Light lecture", [st(1, "Light statement one.", [LEC]), st(2, "Light statement two.", [LEC])], "light")
    return s


def lectures(storage):
    return md.load_lectures([{"title": "First: full", "job": "job_full"}, {"title": "Second: light", "job": "job_light"}], storage)


def text(slide):
    return "\n".join(p.text for sh in slide.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs)


def merged(storage, **kw):
    warnings = []
    prs = Presentation(io.BytesIO(md.merge_deck(lectures(storage), warnings=warnings, **kw)))
    return prs, warnings


def test_mode_is_detected_from_the_files(storage):
    assert md.detect_mode("job_full", storage) == ("full", True)
    assert md.detect_mode("job_light", storage) == ("light", False)
    # a job whose report says light is light even if a notes_final.json exists
    save_json(job_key("job_full", "pipeline_report.json"), {"mode": "light"}, storage)
    assert md.detect_mode("job_full", storage) == ("light", False)


def test_ids_are_prefixed_so_two_jobs_do_not_collide(storage):
    lec = lectures(storage)
    ids = [s.id for l in lec for s in l.statements]
    assert len(ids) == len(set(ids)) == 4 and ids[0] == "L1_st_0001" and ids[2] == "L2_st_0001"
    assert lec[0].sections[0].id == "L1_sec_01" and lec[1].statements[0].section_id == "L2_sec_01" and lec[1].slides[0].section_id == "L2_sec_01"


def test_slide_count_is_title_plus_one_divider_per_lecture_plus_content_plus_sources(storage):
    prs, _ = merged(storage)
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["Lecture notes", "First: full", "Full lecture", "Second: light", "Light lecture", "Sources"]


def test_dividers_show_title_and_mode(storage):
    prs, _ = merged(storage)
    first, second = text(prs.slides[1]), text(prs.slides[3])
    assert "LECTURE 1 OF 2" in first and "full: source-checked" in first and "light: not judged" not in first
    assert "LECTURE 2 OF 2" in second and "light: not judged" in second and "source-checked" not in second


def test_light_statements_are_never_labelled_source_checked(storage):
    prs, _ = merged(storage)
    light_slide = text(prs.slides[4])
    assert "Light statement one." in light_slide and "not judged" in light_slide and "source-checked" not in light_slide
    full_slide = text(prs.slides[2])
    assert "source-checked" in full_slide and "rewritten" in full_slide  # st_0002 was rewritten by the fixer
    # title slide counts only what the judge checked: 2 of 4 statements
    assert "2 sections - 4 statements - 2 source-checked by the judge" in text(prs.slides[0])


def test_timestamps_and_web_links_are_kept(storage):
    prs, _ = merged(storage)
    assert "Lecture 00:12‑00:27" in text(prs.slides[2]) and "Lecture 00:12‑00:27" in text(prs.slides[4])
    links = [r.hyperlink.address for sh in prs.slides[2].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
    assert links == [WEB["url"]]
    sources = text(prs.slides[-1])
    assert "A page (example.org)" in sources and "No web sources were used" not in sources


def test_slide_numbers_are_written_again_after_the_dividers_are_inserted(storage):
    prs, _ = merged(storage)
    for number, slide in enumerate(prs.slides, start=1):
        assert text(slide).strip().splitlines()[-1] == str(number)


def test_nothing_leaves_the_slide(storage):
    prs, _ = merged(storage)
    for slide in prs.slides:
        for sh in slide.shapes:
            assert sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= prs.slide_width and sh.top + sh.height <= prs.slide_height


def test_merged_notes_have_a_header_per_lecture_and_not_judged_only_for_light(storage):
    notes = md.merge_notes(lectures(storage), title="All lectures")
    assert notes.startswith("# All lectures")
    assert "1. First: full (full: source-checked)" in notes and "2. Second: light (light: not judged)" in notes
    full_part, light_part = notes.split("# Lecture 2: Second: light")
    assert "# Lecture 1: First: full" in full_part and "NOT JUDGED" not in full_part
    assert "NOT JUDGED" in light_part and "Light statement one." in light_part and "00:12" in light_part
    assert "source-checked" not in light_part.lower()


def test_a_job_without_slides_is_a_clear_error(storage):
    save_json(job_key("job_empty", "sections.json"), [], storage)
    with pytest.raises(FileNotFoundError, match="slides.json is missing"):
        md.load_lecture(1, "x", "job_empty", storage)


def test_missing_judge_files_never_claim_source_checked(storage):
    make_job(storage, "job_nojudge", "No judge files", [st(1, "A statement.", [LEC])], "full", judge_files=False)
    lec = md.load_lecture(1, "x", "job_nojudge", storage)
    assert lec.mode == "full" and set(lec.status.values()) == {"not judged"}
